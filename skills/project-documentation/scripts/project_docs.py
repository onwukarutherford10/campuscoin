#!/usr/bin/env python3
"""Generate a source-grounded technical report without executing target code."""
from __future__ import annotations
import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from analyze import analyze, files, PRIVATE
from render import render

SCHEMA=Path(__file__).resolve().parents[1]/'schemas/project-documentation.schema.json'


def validate_model(doc,root):
    import jsonschema
    jsonschema.validate(doc,json.loads(SCHEMA.read_text()))
    allowed=set(files(root))
    editorial_sources=[s for blocks in doc.get('editorial',{}).values() for block in blocks for s in block['sources']]
    for source in doc['provenance']+[f['source'] for f in doc['code_figures']]+editorial_sources:
        if source['path'] not in allowed or PRIVATE.search(source['path']):
            raise ValueError('Invalid or private evidence path: '+source['path'])
        if source['end_line']>len((root/source['path']).read_text(errors='replace').splitlines()):
            raise ValueError('Evidence range exceeds source file: '+source['path'])
    serialized=json.dumps(doc)
    if re.search(r'(?i)://[^\s"/@:]+:[^\s"/@]+@',serialized): raise ValueError('Credentialed URL in model')
    if re.search(r'-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----',serialized): raise ValueError('Private key in model')


def validate_pdf(pdf,doc,qa_dir):
    import pymupdf
    pdf=Path(pdf)
    if not pdf.exists() or pdf.stat().st_size<1000: raise ValueError('PDF missing or empty')
    d=pymupdf.open(pdf)
    if len(d)<3: raise ValueError('Expected cover, contents, and body pages')
    text='\n'.join(p.get_text() for p in d)
    expected=['Contents','Project Overview','Technology Stack','System Architecture','Implementation References']
    if doc['data_model']['tables']: expected.append('Database Design')
    if doc['api']['routes']: expected.append('API Design')
    if doc['code_figures']: expected.append('Selected Implementation')
    missing=[x for x in expected if x not in text]
    if missing: raise ValueError('Missing expected sections: '+', '.join(missing))
    if re.search(r'\{\{[^}]+\}\}|\[\[TODO\]\]|-----BEGIN PRIVATE KEY-----',text): raise ValueError('Placeholder or private key in PDF')
    if re.search(r'(?i)://[^\s/@:]+:[^\s/@]+@',text): raise ValueError('Credentialed URL in PDF')
    qa_dir=Path(qa_dir); qa_dir.mkdir(parents=True,exist_ok=True)
    indices={0,1,len(d)-1}
    terms=['System Architecture','Database Design','Selected Implementation','API Design']
    for term in terms:
        for i,p in enumerate(d):
            if i>1 and term in p.get_text(): indices.add(i); break
    for i in sorted(indices): d[i].get_pixmap(matrix=pymupdf.Matrix(1.2,1.2),alpha=False).save(qa_dir/f'page-{i+1:02d}.png')
    result={'page_count':len(d),'expected_sections':expected,'diagram_count':1+len(list((pdf.parent.parent/'diagrams').glob('database-er-*.json'))),'code_figure_count':len(doc['code_figures']),'qa_pages':[i+1 for i in sorted(indices)],'file_size':pdf.stat().st_size}
    (pdf.parent.parent/'validation.json').write_text(json.dumps(result,indent=2))
    return result


def validate_docx(path,model):
    from docx import Document
    path=Path(path)
    if not path.exists() or path.stat().st_size<1000: raise ValueError('DOCX missing or empty')
    document=Document(path)
    text='\n'.join([p.text for p in document.paragraphs]+[cell.text for table in document.tables for row in table.rows for cell in row.cells])
    expected=['Contents','Project Overview','Technology Stack','System Architecture','Implementation References']
    if model['data_model']['tables']: expected.append('Database Design')
    if model['api']['routes']: expected.append('API Design')
    if model['code_figures']: expected.append('Selected Implementation')
    missing=[x for x in expected if x not in text]
    if missing: raise ValueError('Missing DOCX sections: '+', '.join(missing))
    if re.search(r'\{\{[^}]+\}\}|\[\[TODO\]\]|-----BEGIN PRIVATE KEY-----|(?i:gh[pousr]_[A-Za-z0-9_]{20,})|(?i://[^\s/@:]+:[^\s/@]+@)',text):
        raise ValueError('Placeholder or secret pattern in DOCX')
    with zipfile.ZipFile(path) as z:
        images=[x for x in z.namelist() if x.startswith('word/media/')]
        from render import er_specs
        expected_images=1+len(er_specs(model))
        if len(images)<expected_images: raise ValueError('DOCX diagrams missing')
    if len(document.tables)<2: raise ValueError('DOCX tables missing')
    for figure in model['code_figures']:
        sample=next((line.strip() for line in figure['code'].splitlines() if line.strip()),'')
        if sample and sample not in text:
            raise ValueError('Editable code figure text missing: '+figure['path'])
    result={'file_size':path.stat().st_size,'table_count':len(document.tables),'diagram_image_count':len(images),'expected_sections':expected}
    (path.parent.parent/'docx-validation.json').write_text(json.dumps(result,indent=2))
    return result


def main(argv=None):
    ap=argparse.ArgumentParser(prog='project-docs')
    sub=ap.add_subparsers(dest='command',required=True)
    g=sub.add_parser('generate',help='Analyze a repository and generate technical documentation')
    g.add_argument('repository',type=Path)
    g.add_argument('--output-dir',type=Path)
    g.add_argument('--analysis-only',action='store_true')
    g.add_argument('--format',choices=['pdf','docx','both'],default='pdf')
    g.add_argument('--editorial',type=Path,help='Optional source-backed narrative JSON merged into the analysis model')
    args=ap.parse_args(argv)
    root=args.repository.resolve()
    if not root.is_dir(): ap.error('repository must be a directory')
    out=(args.output_dir or root/'.project-docs').resolve()
    if out==root or root in out.parents and out.name not in {'.project-docs','project-docs-build'} and not args.output_dir:
        ap.error('unsafe output directory')
    analysis_dir=out/'analysis'; analysis_dir.mkdir(parents=True,exist_ok=True)
    doc=analyze(root)
    if args.editorial:
        editorial=json.loads(args.editorial.read_text())
        doc['editorial']=editorial
    validate_model(doc,root)
    (analysis_dir/'project-documentation.json').write_text(json.dumps(doc,indent=2,ensure_ascii=False))
    for key,name in [('metadata','repository'),('architecture','architecture'),('data_model','database'),('api','api'),('provenance','provenance')]:
        (analysis_dir/f'{name}.json').write_text(json.dumps(doc[key],indent=2,ensure_ascii=False))
    if args.analysis_only: print(analysis_dir/'project-documentation.json'); return 0
    result={}
    if args.format in {'pdf','both'}:
        pdf=render(doc,out/'output'/f"{doc['metadata']['name']}-technical-documentation.pdf",out/'diagrams')
        result['pdf']=str(pdf)
        result['pdf_validation']=validate_pdf(pdf,doc,out/'qa')
    if args.format in {'docx','both'}:
        from render_docx import render_docx
        docx=render_docx(doc,out/'output'/f"{doc['metadata']['name']}-technical-documentation.docx",out/'figures')
        result['docx']=str(docx)
        result['docx_validation']=validate_docx(docx,doc)
    print(json.dumps(result,indent=2))
    return 0


if __name__=='__main__': sys.exit(main())
