"""Editable DOCX composition from the same ProjectDocumentation JSON as the PDF."""
from __future__ import annotations
import io
import json
import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfgen import canvas
import pymupdf
from pygments import lex
from pygments.lexers import get_lexer_by_name
from pygments.token import Token
from render import architecture_spec, er_specs, Diagram, source_text

NAVY=RGBColor(23,34,53)
TEAL=RGBColor(30,86,122)
MUTED=RGBColor(83,96,112)
WHITE=RGBColor(255,255,255)
DOC_WIDTH=6.75


def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),color)


def set_cell_text(cell, value, bold=False):
    cell.text=''
    p=cell.paragraphs[0]
    p.style='Normal'
    r=p.add_run(str(value))
    r.bold=bold; r.font.name='Arial'; r.font.size=Pt(8)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER


def keep_row_together(row):
    trPr=row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def add_table(doc, rows, widths=None):
    t=doc.add_table(rows=1,cols=len(rows[0]))
    t.autofit=False
    if widths:
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    t.style='Table Grid'
    for i,v in enumerate(rows[0]):
        if widths: t.rows[0].cells[i].width=Inches(widths[i])
        set_cell_text(t.rows[0].cells[i],v,True); shade(t.rows[0].cells[i],'EAF0F3')
    keep_row_together(t.rows[0])
    for row in rows[1:]:
        new_row=t.add_row();keep_row_together(new_row);cells=new_row.cells
        for i,v in enumerate(row):
            if widths: cells[i].width=Inches(widths[i])
            set_cell_text(cells[i],v)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t


def caption(doc,text):
    p=doc.add_paragraph(style='Caption');p.add_run(text)


def body(doc,text):
    return doc.add_paragraph(str(text),style='Normal')


def heading(doc,text,level=1):
    return doc.add_heading(str(text),level=level)


def render_diagram(spec,path,width=540,height=300):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    mem=io.BytesIO();c=canvas.Canvas(mem,pagesize=(width,height))
    Diagram(spec,width,height).drawOn(c,0,0);c.showPage();c.save()
    pdf=pymupdf.open(stream=mem.getvalue(),filetype='pdf')
    page=pdf[0]
    page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False).save(path)
    return path


def add_diagram(doc,spec,path,caption_text,height=300):
    render_diagram(spec,path,height=height)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(DOC_WIDTH))
    caption(doc,caption_text)


def code_figure(doc,fig,number):
    heading(doc,fig['caption'],2)
    explanation=body(doc,fig['explanation'])
    explanation.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=1)
    keep_row_together(t.rows[0])
    t.autofit=False;t.columns[0].width=Inches(DOC_WIDTH)
    cell=t.cell(0,0);shade(cell,'111820');cell.text=''
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4)
    r=p.add_run(fig['path']);r.font.name='Consolas';r.font.size=Pt(8);r.font.color.rgb=RGBColor(184,199,209)
    try:lexer=get_lexer_by_name({'tsx':'typescript','ts':'typescript','py':'python','js':'javascript','jsx':'javascript'}.get(fig['language'],fig['language']))
    except Exception:lexer=None
    for i,line in enumerate(fig['code'].splitlines()):
        p=cell.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
        r=p.add_run(f"{fig['start_line']+i:>3}  ");r.font.name='Consolas';r.font.size=Pt(7);r.font.color.rgb=RGBColor(114,133,150)
        shown=line.expandtabs(2)
        # Keep editable text; Word and Docs may wrap long source lines within the panel.
        parts=list(lex(shown,lexer)) if lexer else [(Token.Text,shown)]
        for token,value in parts:
            if '\n' in value:value=value.split('\n')[0]
            if not value:continue
            rgb=RGBColor(224,232,238)
            if token in Token.Keyword:rgb=RGBColor(215,165,233)
            elif token in Token.String:rgb=RGBColor(157,206,164)
            elif token in Token.Comment:rgb=RGBColor(132,150,165)
            elif token in Token.Name.Function or token in Token.Name.Class:rgb=RGBColor(133,199,233)
            elif token in Token.Number:rgb=RGBColor(233,186,131)
            r=p.add_run(value);r.font.name='Consolas';r.font.size=Pt(7);r.font.color.rgb=rgb
    caption(doc,f"Figure {number}. Lines {fig['start_line']}-{fig['end_line']} of {fig['path']}.")


def configure(doc):
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.left_margin=Inches(.88);sec.right_margin=Inches(.88)
    sec.top_margin=Inches(.82);sec.bottom_margin=Inches(.72)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10.5);normal.font.color.rgb=NAVY
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
    title=doc.styles['Title'];title.font.name='Times New Roman';title.font.size=Pt(26);title.font.bold=True;title.font.color.rgb=NAVY
    title.paragraph_format.space_after=Pt(12)
    title_ppr=title._element.get_or_add_pPr()
    title_border=title_ppr.find(qn('w:pBdr'))
    if title_border is not None: title_ppr.remove(title_border)
    for name,size in [('Heading 1',17),('Heading 2',12)]:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=NAVY
        st.paragraph_format.space_before=Pt(15);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
    cap=doc.styles['Caption'];cap.font.name='Times New Roman';cap.font.size=Pt(9);cap.font.italic=True;cap.font.color.rgb=MUTED
    cap.paragraph_format.space_after=Pt(8)
    header=sec.header.paragraphs[0];header.text='';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer=sec.footer.paragraphs[0];footer.text='';footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT


def render_docx(model,path,figures_dir):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    figures_dir=Path(figures_dir);figures_dir.mkdir(parents=True,exist_ok=True)
    d=Document();configure(d)
    md=model['metadata'];title=md.get('display_name') or md['name'].replace('-',' ').title()
    d.core_properties.title=title+' Technical Documentation';d.core_properties.subject='Repository architecture and implementation'
    p=d.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(title)
    p=d.add_paragraph('Repository architecture and implementation report');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(f"Branch: {md.get('branch') or 'not available'}\nCommit: {(md.get('commit') or 'not available')[:12]}\nGenerated: {md['generated_at'][:10]}");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    heading(d,'Contents')
    sections=['Project Overview','Technology Stack','System Architecture','Component Architecture']
    if model['data_model']['tables']:sections.append('Database Design')
    if model['api']['routes']:sections.append('API Design')
    if model.get('security'):sections.append('Authentication and Security')
    if model.get('background_processing') or model.get('integrations'):sections.append('Operations and Integrations')
    if model['testing']['files']:sections.append('Testing Strategy')
    if model['code_figures']:sections.append('Selected Implementation')
    if model['limitations']:sections.append('Limits of Repository Evidence')
    sections.append('Implementation References')
    numbered={name:f'{i}. {name}' for i,name in enumerate(sections[:-1],1)}
    for name in sections:body(d,numbered.get(name,name))
    body(d,'In Google Docs, use Insert > Table of contents after import to add live page numbers.')
    d.add_page_break()
    heading(d,numbered['Project Overview'])
    body(d,f"{title} is classified by repository structure as a {md['type']} project. This report documents discovered source code, manifests, routes, persistence declarations, and tests at the recorded commit. Classification is a structural inference rather than a runtime deployment claim.")
    caption(d,'Evidence: '+', '.join(md['manifests'][:4]))
    heading(d,numbered['Technology Stack'])
    rows=[['Technology','Repository evidence']]
    rows += [[f,', '.join(md['manifests'][:3])] for f in md['frameworks']]
    rows += [[x['name'],f"{x['files']} source files"] for x in md['languages'][:6]]
    if len(rows)>1:add_table(d,rows,[2.4,4.35])
    heading(d,numbered['System Architecture'])
    body(d,'The component view reflects source directories and declared integration points. Arrows indicate likely call direction across recognized layers; individual flows should be checked against the referenced modules.')
    add_diagram(d,architecture_spec(model),figures_dir/'architecture.png','Figure 1. Repository component architecture.',height=max(210,55*len(model['architecture']['nodes'])))
    heading(d,numbered['Component Architecture'])
    add_table(d,[['Component','Files','Example evidence']]+[[g['name'],str(g['count']),source_text(g['sources'][0])] for g in model['architecture']['components']],[2.25,.6,3.9])
    db=model['data_model'];ers=er_specs(model)
    diagrams_dir=figures_dir.parent/'diagrams';diagrams_dir.mkdir(parents=True,exist_ok=True)
    (diagrams_dir/'architecture.json').write_text(json.dumps(architecture_spec(model),indent=2))
    for i,spec in enumerate(ers,1):
        (diagrams_dir/f'database-er-{i}.json').write_text(json.dumps(spec,indent=2))
    if db['tables']:
        heading(d,numbered['Database Design'])
        if db.get('dialect'):body(d,f"Configured database dialect: {db['dialect']}. Evidence: {source_text(db['dialect_source'])}.")
        body(d,f"Static analysis found {len(db['tables'])} {db.get('technology') or 'schema'} model declarations. Migration operations are cross-referenced where present. Each ER panel shows selected fields and visible relationships.")
        for i,spec in enumerate(ers,1):
            heading(d,f'Entity relationship panel {i}',2)
            add_diagram(d,spec,figures_dir/f'database-er-{i}.png',f'Figure {i+1}. Entity relationships, panel {i}.',height=((len(spec['nodes'])+1)//2)*112+30)
            rel=', '.join(f"{x['from']}.{x['label']} to {x['to']}" for x in spec['edges']) or 'No foreign keys are visible within this panel.'
            body(d,'Visible relationships: '+rel)
        for t in db['tables']:
            heading(d,t['name'],2)
            body(d,f"ORM class {t['class']}. Model: {source_text(t['source'])}. "+(f"Migration: {source_text(t['migration_sources'][0])}." if t['migration_sources'] else 'No matching create-table migration found in static analysis.'))
            rows=[['Column','Type / relationship']]+[[c['name'],c['type']+(' to '+c['foreign_key'] if c.get('foreign_key') else '')+(' (PK)' if c.get('primary_key') else '')] for c in t['columns']]
            if len(rows)>1:add_table(d,rows,[2.4,4.35])
    routes=model['api']['routes']
    if routes:
        heading(d,numbered['API Design'])
        body(d,f"Static decorator analysis identified {len(routes)} route handlers. Authentication labels reflect decorators present directly on handlers. Named request schemas are listed where explicit.")
        for group in sorted({r['group'] for r in routes}):
            heading(d,group.title(),2)
            add_table(d,[['Method','Route','Auth','Source']]+[[r['method'],r['path'],r['auth'],source_text(r['source'])] for r in routes if r['group']==group],[.65,2.55,1.25,2.3])
        validated=[r for r in routes if r.get('request_schema')]
        if validated:
            heading(d,'Declared request validation',2)
            add_table(d,[['Method and route','Schema','Source']]+[[r['method']+' '+r['path'],r['request_schema'],source_text(r['source'])] for r in validated],[3.1,1.6,2.05])
    if model.get('security'):
        heading(d,numbered['Authentication and Security'])
        body(d,'These controls were identified from executable source patterns. This list does not represent a complete security audit.')
        add_table(d,[['Control','Implementation evidence']]+[[x['name'],source_text(x['source'])] for x in model['security']],[3.4,3.35])
    if model.get('background_processing') or model.get('integrations'):
        heading(d,numbered['Operations and Integrations'])
        body(d,'Commands and integration boundaries below are source declarations. Production scheduling or availability is not established by these declarations alone.')
        if model.get('background_processing'):add_table(d,[['Command','Source']]+[[x['name'],source_text(x['source'])] for x in model['background_processing']],[3.4,3.35])
        if model.get('integrations'):add_table(d,[['Integration','Source']]+[[x['name'],source_text(x['source'])] for x in model['integrations']],[3.4,3.35])
    tests=model['testing']['files']
    if tests:
        heading(d,numbered['Testing Strategy'])
        body(d,f'The repository contains {len(tests)} discovered test files. This indicates test presence, not measured coverage or passing status.')
        from collections import Counter
        counts=Counter('/'.join(p.split('/')[:3]) if len(p.split('/'))>3 else '/'.join(p.split('/')[:2]) for p in tests)
        add_table(d,[['Test area','Files']]+[[k,str(v)] for k,v in sorted(counts.items())],[5.6,1.15])
    if model['code_figures']:
        heading(d,numbered['Selected Implementation'])
        for i,f in enumerate(model['code_figures'],1):code_figure(d,f,i+len(ers)+1)
    if model['limitations']:
        heading(d,numbered['Limits of Repository Evidence'])
        for text in model['limitations']:body(d,'• '+text)
    heading(d,'Implementation References')
    body(d,'Paths are relative to the repository root. The JSON analysis artifact retains line ranges for every machine-generated claim.')
    add_table(d,[['Evidence path','Line']]+[[s['path'],str(s['start_line'])] for s in model['provenance'][:8]],[5.5,1.25])
    d.save(path)
    return path
